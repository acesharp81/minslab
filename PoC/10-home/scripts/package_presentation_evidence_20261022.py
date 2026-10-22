"""Build shareable, source-scoped presentation evidence from read-only extracts."""
import csv,json,hashlib,shutil,re,subprocess,zipfile,os
from collections import Counter,defaultdict
from datetime import datetime,date,timezone
from pathlib import Path
from urllib.parse import urlsplit
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment
from backend.policy_delivery_comparison import _events,_forecast_points
from scripts.extract_presentation_evidence_20261022 import clean,csvwrite as _unused

ROOT=Path(__file__).resolve().parents[1];A=ROOT/'docs/analysis';OUT=A/'THE_JUTAEK_PRESENTATION_EVIDENCE_20261022';DAY=date(2026,10,22)
STAGES={'POLICY':'정책 수립','BUSINESS':'사업화','APPROVAL':'인허가','CONSTRUCTION':'착공','COMPLETION':'준공','SUPPLY':'공급·모집','MOVE_IN':'입주'}

def read(path):return json.loads((OUT/path).read_text())
def old(path):return json.loads((ROOT/path).read_text())
def rows(path):
    csv.field_size_limit(10000000)
    with path.open(encoding='utf-8-sig') as f:return list(csv.DictReader(f))
def dump(path,data):
    p=OUT/path;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(clean(data),ensure_ascii=False,indent=2,default=str))
def md(path,text):
    p=OUT/path;p.parent.mkdir(parents=True,exist_ok=True);p.write_text(clean(text.strip())+'\n')
def csvout(path,records,fields=None):
    records=list(records);fields=fields or list(dict.fromkeys(k for r in records for k in r)) or ['status'];p=OUT/path;p.parent.mkdir(parents=True,exist_ok=True)
    with p.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore');w.writeheader()
        for row in records:
            row=clean(row);record={k:json.dumps(v,ensure_ascii=False,default=str) if isinstance(v,(dict,list)) else v for k,v in row.items()}
            w.writerow({k:"'"+v if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@')) else v for k,v in record.items()})
def copy_text(src,dst):md(dst,(ROOT/src).read_text(encoding='utf-8-sig'))
def table(headers,data):
    esc=lambda v:str('미확보' if v is None else v).replace('|',' / ').replace('\n',' ')
    return '\n'.join(['| '+' | '.join(headers)+' |','|'+'|'.join(['---']*len(headers))+'|']+['| '+' | '.join(esc(v) for v in r)+' |' for r in data])
def workbook(path,sheets):
    w=Workbook();w.remove(w.active)
    for name,records in sheets:
        s=w.create_sheet(name[:31]);records=list(records);keys=list(dict.fromkeys(k for r in records for k in r)) or ['status'];s.append(keys)
        for row in records:
            vals=[]
            for key in keys:
                v=clean(row.get(key))
                if isinstance(v,(dict,list)):v=json.dumps(v,ensure_ascii=False,default=str)
                if isinstance(v,str):
                    v=re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]','',v)
                    if v.lstrip().startswith(('=','+','-','@')):v="'"+v
                    if len(v)>32000:v=v[:31900]+' [엑셀 셀 제한으로 생략; CSV/JSON 원본 참조]'
                vals.append(v)
            s.append(vals)
        for cell in s[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='194F7E');cell.alignment=Alignment(wrap_text=True,vertical='center')
        s.freeze_panes='A2';s.auto_filter.ref=s.dimensions;s.row_dimensions[1].height=34
        for col in s.columns:
            letter=col[0].column_letter;s.column_dimensions[letter].width=min(48,max(14,len(str(col[0].value or ''))*1.6))
        for row in s.iter_rows(min_row=2):
            for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
    p=OUT/path;p.parent.mkdir(parents=True,exist_ok=True);w.save(p)

def identity_and_map(model,projects,programs):
    cohorts=model['portfolio_cohorts'];index={p['id']:p for p in projects};cohort_member={pid:c['id'] for c in cohorts for pid in c['source_ids']}
    maps=[];events=[];targets=[];summary=[];ledger=[]
    for stage,label in list(STAGES.items())[1:]:
        actual=_events(cohorts,stage,DAY,'1900-01-01',cumulative=True);future=_forecast_points(cohorts,stage,DAY);f=next((p for p in reversed(future) if p['date']<='2026-12-31'),{})
        summary.append(dict(stage=stage,label=label,as_of_date=str(DAY),current_units=actual[-1]['units'] if actual else None,year_end_units=f.get('units'),remaining_plan_units=f.get('remaining_plan_units'),unit='호',project_count=len(cohorts),rule='동일 관리사업 최대 도달량·총량 상한 적용; 1900년부터 기준일까지 누적. 후속 단계 포함. 단계 간 합산 금지'))
    for c in cohorts:
        p=index[c['id']];m=p.get('mapping',{});assessment=p.get('supply_type_assessment',{})
        maps.append(dict(project_id=c['id'],project_name=c['name'],policy_id='POL_2025_0907',program_id=c.get('program_id'),region_code=c.get('municipality_id'),region=c.get('region'),municipality=c.get('municipality'),address_or_area_name=p.get('address') or ' '.join([c.get('region') or '',c.get('municipality') or '',p.get('district') or '',c.get('block') or '']),latitude=c.get('latitude'),longitude=c.get('longitude'),coordinate_basis=c.get('coordinate_basis'),geometry_availability='좌표 보유' if c.get('latitude') is not None and c.get('longitude') is not None else '행정구역만 보유' if c.get('municipality_id') else '위치 미확보',supply_type=assessment.get('label'),supply_form=p.get('tenure'),planned_units=c.get('planned_units'),adopted_capacity=c['units'],current_stage=c.get('representative_stage'),stage_status=c.get('eligibility_status'),last_updated_at=p.get('recorded_at'),source_evidence_id=c['source_ids'],mapping_status=m.get('status'),evidence_url=c.get('evidence_url')))
        for stage,label in list(STAGES.items())[1:]:
            actual=_events([c],stage,DAY,'1900-01-01',cumulative=True);cur=actual[-1]['units'] if actual else None
            forecast=_forecast_points([c],stage,DAY);last=next((p for p in reversed(forecast) if p['date']<='2026-12-31'),{})
            ledger.append(dict(project_id=c['id'],project_name=c['name'],stage=stage,units=cur,capacity=c['units'],year_end_units=last.get('units'),source_ids=c['source_ids'],as_of_date=str(DAY),metric='누적 도달량·계획',contains_manual=any(e.get('origin')=='MANUAL' for e in c['events'] if e['stage']==stage),contains_estimates=any(e.get('inferred') for e in c['events'] if e['stage']==stage)))
        for e in c['events']:
            events.append(dict(project_id=c['id'],project_name=c['name'],source_project_id=e.get('project_id'),milestone_type=e['stage'],milestone_label=STAGES.get(e['stage'],e['stage']),initial_planned_date=None,revised_planned_date=None,actual_or_observed_date=e.get('event_date'),date_basis=e.get('date_basis'),status=e.get('status'),planned_units=None,actual_units=e.get('units'),changed_at=e.get('changed_at'),change_reason=e.get('input_basis') or e.get('basis'),source_evidence_id=e.get('source_entity_id'),evidence_url=e.get('evidence_url'),origin=e.get('origin'),review_status=e.get('review_status'),inferred=e.get('inferred'),evidence_grade='시스템/담당자 추정' if e.get('inferred') else '업무담당자 확인' if e.get('origin')=='MANUAL' else '수집 원천',source_fact_grade='공식/비공식은 원문 URL·근거를 별도 확인',event_record_is_not_total=True))
        for p in c['schedules']:
            targets.append(dict(project_id=c['id'],project_name=c['name'],milestone_type=p['stage'],initial_planned_date=(p.get('baseline') or {}).get('due_date'),revised_planned_date=p.get('due_date'),actual_date=None,status='PLAN',planned_units=p.get('units'),actual_units=None,changed_at=p.get('changed_at'),change_reason=p.get('input_basis') or p.get('basis'),source_evidence_id=p.get('source_entity_id'),evidence_url=p.get('evidence_url'),origin=p.get('origin'),inferred=p.get('inferred'),date_precision=p.get('precision'),date_label=p.get('source_date_label'),note='baseline은 최초 원문 계획이 아닌 담당자 채택 기준선일 수 있음. 날짜 없음은 null 유지.'))
    csvout('09_map_dashboard/project_map.csv',maps);csvout('09_map_dashboard/project_milestones.csv',events);csvout('09_map_dashboard/progress_targets.csv',targets);csvout('09_map_dashboard/stage_summary.csv',summary);csvout('09_map_dashboard/project_stage_counting_ledger.csv',ledger)
    csvout('04_entity_linking/entity_dictionary.csv',[dict(entity_type='POLICY',entity_id=p.get('policy_id'),name=p['name'],scope=p.get('target_scope'),status=p.get('lifecycle_status')) for p in read('00_data/policies.json')]+[dict(entity_type='PROGRAM',entity_id=p['id'],name=p['name'],parent_id=p.get('policy_id'),status=p.get('status')) for p in programs]+[dict(entity_type='SOURCE_PROJECT',entity_id=p['id'],name=p['name'],parent_id=p.get('mapping',{}).get('program_id'),source_id=p.get('source_id'),district=p.get('district'),block=p.get('block'),source_object_id=p.get('source_object_id'),status=p.get('mapping',{}).get('status'),management_id=cohort_member.get(p['id'])) for p in projects])
    return summary,maps,events,targets,ledger

CROSS={'M01':'R001','M02':'R011','M03':'R013','M04':'R014','M05':'R016','M06':'R020','M07':'R021','M08':'R023','M10':'R033','M11':'R034','M14':'R038','M15':'R046','M16':'R051','M17':'R052','M19':'R059','M20':'R060','M23':'R070','M24':'R071','M25':'R072','M30':'R073','M31':'R074','S01':'R024','S02':'R025','S03':'R026','S04':'R027','S05':'R043','S06':'R042','S07':'R044','S13':'R067','S17':'R028','S18':'R029','S19':'R030'}

def official_validation(summary):
    detailed=rows(A/'reliability_update_20261022/official_185_comparison.csv');by={r['지표ID']:r for r in detailed};base=old('backend/data/official_benchmark_20261016.json');oldagreement=old('docs/analysis/reliability_update_20261022/round.json')['agreement'];legacy={r['id']:r for r in oldagreement['rows']}
    additional=rows(A/'official_followup_sources_20261021/followup_results.csv');additional_by={r['check_id']:r for r in additional}
    quantity=rows(A/'reliability_update_20261022/official-32.csv');starts=rows(A/'gpt_report_20261020/comparison_start_8.csv');statistics=rows(A/'official_stat_verification_20261020/verified_comparison.csv')
    for key,needle,value in [('M32','인허가','84'),('M33','착공','87')]:
        hit=next((r for r in detailed if needle in r['지표'] and str(r['공식 기준수치']) in (value,value+'.0') and r['단위']=='%'),None)
        if hit:CROSS[key]=hit['지표ID']
    roster=[]
    for b in base['rows']+base['supplemental_site_rows']:
        link=by.get(CROSS.get(b['id']),{});prior=legacy.get(b['id'],{});home=link.get('HOME_동일범위비교값') or None;status=link.get('판정') or '판정 미산출';note=link.get('사유') or '이 56목록은 기존53+보충3의 재구성. 외부 요청의 원래56 정의와 동일한지 미확인.'
        # Keep the government site examples' quantity and timing qualifications distinct.
        if b['id'] in ('S01','S02','S03','S04','S05','S06','S07','S17'):
            lookup={'S01':'고양창릉','S02':'인천계양','S03':'하남교산','S04':'남양주왕숙2','S05':'중계본동','S06':'흑석11','S07':'신림2','S17':'부천대장'}[b['id']]
            current=next(r for r in starts if r['project']==lookup);home=current['home_comparison_units'];status=current['assessment'];note=current['reason']+'; '+current.get('verification','')+'; 공식 구성블록·사건기간의 전량 검증 여부는 별도.'
        if b['id'] in ('M27','M29','M37','S14','S15','S16'):status='비수량/계획 확인 항목';home=None
        aid={'M16':'A01','M19':'A02','M32':'A03','M33':'A04','M22':'A05','S14':'A06','S15':'A08','S16':'A08'}.get(b['id']);extra=additional_by.get(aid,{})
        if extra:
            home=extra.get('home_value') or None;status=extra['match_status'];note=extra['difference_reason']+'; '+extra.get('unresolved_requirement','')
            if b['id']=='S15':home='창릉 S6 2027-12'
            if b['id']=='S16':home='부천대장 A5·A6 2027-11'
        roster.append(dict(metric_id=b['id'],명단구분='기존 53개 기준표' if b in base['rows'] else '별도 보충 3개',원문페이지_표=(link.get('원문 내 위치') or b.get('official_url','').split('#')[-1] if '#' in b.get('official_url','') else link.get('원문 내 위치')),원문지표명=b['name'],상위주제=b['kind'],하위구분=b.get('stage'),공식수치=b['official_value'],공식상한=b.get('official_upper'),프로젝트산출수치=home,단위=b['unit'],공식기준일=b.get('official_period'),프로젝트데이터기준일='최신 보관 대조표 2026-10-22 / 개별 근거는 원 작성시점 유지',모집단_정책범위=b.get('official_period'),필터조건=link.get('비교기준'),집계산식=link.get('집계·해석 주의'),비교방법='개략표현·범위·사건기간 분리',일치여부=status,차이=link.get('수량차이'),원본근거URL=b['official_url'],산출근거_쿼리='185_all.csv#'+link.get('지표ID','미대응'),최신185항목ID=link.get('지표ID'),추가검증ID=aid,부분참고값=extra.get('partial_home_units'),증감률기준호수2025=extra.get('home_base_units_2025'),증감률비교호수2026=extra.get('home_comparison_units_2026'),추가증빙URL=extra.get('evidence_urls'),과거53판정=prior.get('status'),과거53값=prior.get('actual_value'),비고=note,독립검증항목여부='아니오 — 총계/부분합/계획/정성항목 포함'))
    assert len(roster)==56 and len({r['metric_id'] for r in roster})==56
    csvout('07_official_validation/56_metrics_comparison.csv',roster);csvout('07_official_validation/185_all.csv',detailed);csvout('07_official_validation/32_q4_projects.csv',quantity);csvout('07_official_validation/8_start_cases.csv',starts);csvout('07_official_validation/10_statistical_checks.csv',statistics);csvout('07_official_validation/8_additional_checks.csv',additional)
    info=[dict(item='56개 목록의 근거',value='기존 공식 기준표53행 + 별도 보충3행. 원 요청자가 정한56개와 동일성은 미확인.'),dict(item='100% 일치 여부',value='증명되지 않음. 56개 전부 일치 주장 금지. 계획·부분합·정성 항목을 독립 지표로 합산하지 않음.'),dict(item='현재185개 판정',value=Counter(r['판정'] for r in detailed)),dict(item='4분기32행',value='32/32 사업 대응; 하나의 사업을 나눈 A1/A3 중복 제거 후31개 수량 비교단위 중26개 원수량 동등.'),dict(item='과거53 평가',value={k:v for k,v in oldagreement.items() if k!='rows'}),dict(item='자료 성격',value='이 파일은 검증기록 추출이다. 미판정 항목을 이번 작업에서 새로 검증 완료로 바꾸지 않았다.')]
    workbook('07_official_validation/07_56_metrics_comparison.xlsx',[('읽기전_중요',info),('56목록_재구성',roster),('최신185항목',detailed),('4분기32행',quantity),('착공8사례',starts),('지역통계10개',statistics),('추가검증8개',additional)])
    official_url='https://www.korea.kr/briefing/policyBriefingView.do?newsId=156785506';appendix='https://www.fsc.go.kr/comm/getFile?srvcId=BBSTY1&upperNo=87948&fileTy=ATTACH&fileNo=3'
    source=[dict(source_id='G01',title='제1차 주택공급 정례브리핑',published='2026-10-16',url=official_url,location='주택공급 디지털 상황판 구축·운영 계획 및 수도권 착공 설명',status='공식 원문 재접속 확인 2026-10-22'),dict(source_id='G04',title='주택공급 계획부터 성과까지 정례적으로 공개',published='2026-10-16',url='https://www.fsc.go.kr/po010101/87948',location='본문·별첨1 4~6쪽',status='공식 원문 재접속 확인 2026-10-22'),dict(source_id='G04-PDF',title='제1차 주택공급 정책 브리프 별첨1',published='2026-10-16',url=appendix,location='4쪽 착공사례; 5~6쪽 모집계획',status='기존 원문 보관 및 해시 대조자료 보유')]
    csvout('07_official_validation/official_sources.csv',source)
    main_start=next(r for r in summary if r['stage']=='CONSTRUCTION')['current_units']
    md('07_official_validation/reconciliation_notes.md',f'''
# 공식 수치와 the주택 누적값: 발표에서 사용할 비교 기준

검토일: 2026-10-22. 최신 누적값은 이번 읽기 전용 추출의 177개 관리사업 원장에서 계산했다.

## 1. 요청서의 두 전제 확인

- **56개 전부 일치는 입증되지 않았다.** 코드에 있는 기존53개 기준표와 별도 보충3개를 합친56개 명단을 제공한다. 기존 기준표 자체도 보충3개를 점수에 더하지 않는다고 명시한다. 이 재구성56개가 요청자가 말한56개와 같은 명단인지 확인이 필요하다.
- 최신 별도 대조표는185행이며, 그중41행 `부합`, 1행 `계획부합`, 나머지는 범위 차이·일부범위·공표판본 차이·확인 유보 등이다. 항목 사이에 총계/부분합/계획/정성 정보가 섞여 있어 전체 정확도 백분율로 바꾸지 않는다.
- 과거53개 고정 평가에서는 숫자항목47개 중 확인된2개, 비교가능2개 중2개 부합이 기록됐다. **비교가능2/2=100%를56/56 또는 전체 데이터 정확도100%로 확대하면 안 된다.** 과거 평가는 최신 추가 검증 성과를 전부 반영한 평가도 아니다.

## 2. 판본 차이

185표는 보관된 전체 판정표이며, 재구성56표에는 이후 보강된 착공8사례·추가검증A01~A08의 결과를 별도 ID로 표시했다. 예컨대 R025의 현재 관리비교2188과 인천계양 전체공사구성2342는 임대154호의 범위를 다르게 표시한 값이다. 서로 다른 범위를 같은 분모의 최신 실적으로 합치지 않는다. 서울일반주택착공은 추가검증A04의2229→4165와86.855%를 사용하며, 같은표의인허가+84%는미확보다. A08의2027년블록입주계획은10.16브리핑의세부값으로인용하지않고별도출처를유지한다.

## 3. 99,800호의 근거

2026-10-22 다시 확인한 [정부 브리핑]({official_url})과 [금융위 공동 보도자료](https://www.fsc.go.kr/po010101/87948)는 **9.9만호 이상**으로 표현한다. 검토한 기준표도99,000을 개략 하한으로 저장한다. **99,800이라는 정확 총계의 산출명세는 이번 근거에서 확인하지 못했다.**

요청 숫자끼리의 산술 `99,800 − {main_start:,} = {99800-main_start:,}`는 맞다. 그러나 앞 수치의 원문 근거와 두 수치의 같은 모집단이 확인되지 않았으므로 이것을 확정 미수집량이라고 부르지 않는다. 공식 개략 하한99,000과의 산술 차이는{99000-main_start:,}이며, 이 역시 동일 범위의 누락량을 뜻하지 않는다.

## 3. 서로 다른 두 모집단

| 항목 | 정부 발표 | the주택 화면 누적 |
|---|---|---|
|기간|2026년 실적, 명목상10월 둘째 주까지; 부문별 마감시점 혼합|2026-10-22까지 정책 연결사업의 누적 도달|
|대상|수도권 전체 공급유형 및 통합 후속대책 범위|공개 근거·담당자 확인으로 연결한9·7 관리사업|
|이전 연도 착공|정부가 어떤 사업을 올해 집계에 산입했는지 세부명세 미확보|2024·2025년 착공도 정책 누적에 포함|
|포함 근거|정부 부문별 통합 원장|원천 사건·담당자 확인·선행단계 추정 포함|
|전량 사업명세|발표에서 공개한 일부 지구·사례 외 전체 명세 미확보|관리177사업의 원천ID·사건·집계표 제공|

따라서 {main_start:,}호를 정부99,000호 이상의 직접 재현값이나 수집률로 해석하지 않는다.

## 4. 수치가 재현되는 부분과 남은 격차

- 4분기32행은 **사업명 대응32/32**. A1/A3 묶음 중복 제거 후31개 수량비교단위 중26개 원수량 동등. 실제 공급 완료 실적과 구분한다.
- 공식주택 계획 합14,731호, 관리주택 규모14,753호. +22호는 −54(병점 오피스텔 범위) +50(대장A2) +5(신길B1) −1(A6) +22(A1/A3)의 상쇄다. 뒤4개 차이는 발생 원인이 미확정이며 반올림이라고 임의 처리하지 않는다.
- 검단 AA19 국민991+영구332=1,323호처럼 유형을 나눠 공개 원문을 추적한 사례는 발표에 사용 가능하다. 지구 한 건과 두 유형 원천을 동시에 합산하지 않는다.
- 같은 공식 통계의 수집·재계산10개는 해당 통계가 시스템에 재현됐다는 근거다. 독립적인9·7 개별사업 전수정확도 검증10개로 표현하지 않는다.
- 공공주택·신축매입·민간정비·기타 일반민간·공공지원민간임대의 정부 동일범위 전체 사업명세는 아직 미확보다. 미공개인지, 다른 기관에서 공개 중인데 미수집인지, 공개 문서는 있지만 정책 귀속을 확정 못 한 것인지를 사업별로 일괄 구분할 수 없다.

## 5. 차이 원인별 현재 판정

'''+table(['원인 후보','현재 확인 수준','발표에서 허용되는 표현'],[
['시차만의 문제','전체 격차 원인으로 미확인','기준시점이 다르다. 전부 시차 때문이라고 확정하지 않는다.'],
['미공개 원장','검토한 공개 발표에 전량 사업명세 미포함','미확보 원장 때문에 완전대사 불가. 세상 어디에도 공개되지 않았다고 단정하지 않는다.'],
['공개 원천 미수집','일부 상세 모집범위·기관별 명세의 미완결 기록 보유','해당 경로·기관·부분합의 확보 상황을 표시한다.'],
['정책 연결 유보','근거 불충분·다른 후속정책·중첩 우려 후보를 보류한 기록 보유','담당기관의 정책·사업 공통키와 산입원장이 있으면 확정 가능성이 높아진다.'],
['중복·범위 차이','부평 중복350호, 주택/오피스텔54실 등 실제 설명 사례 확인','확인한 사례의 전후 수치만 제시한다.'],
['이전 착공의 누적','현재 원장에 과거 착공을 포함한 사실 확인','정부의 올해 실적과 같은 기준이라고 말하지 않는다.']])+f'''

## 6. 안전하게 사용할 발표 문장

“공개된 자료를 정책·프로그램·개별 사업으로 연결하고 원문과 판본까지 추적했습니다. 4분기 발표32항목은 전부 관리사업과 대응했고, 유형 분리와 중복 제외를 통해 세부 구성을 설명할 수 있습니다. 다만 정책 누적과 정부의 연간 전체 실적은 범위가 달라 숫자를 억지로 맞추지 않았습니다. 미확인 부분도 원장과 유보 사유로 남겨 다음 자료가 들어오면 다시 대조할 수 있습니다.”

재현: `09_map_dashboard/project_stage_counting_ledger.csv`의 `CONSTRUCTION` 호수를 관리사업별1회 합산. 코드 `_events(..., cumulative=True)` / `backend/policy_delivery_comparison.py`. 연말 예상은 `_forecast_points`로 현재누적과 미도달 계획의 합집합을 계산한다.
''')
    return roster,detailed,quantity,starts,statistics

def process_and_pipeline(capture,sources,programs,projects,summary):
    from backend.project_collection import CATALOG
    registry=rows(OUT/'01_process_sources/source_registry.csv')
    for r in registry:r['수집대상단계']=' | '.join(STAGES.get(k,k) for k,ids in CATALOG.items() if r['source_id'] in ids) or '공통 지역·지도·시장 또는 별도 참고'
    csvout('01_process_sources/source_registry.csv',registry)
    org={s['id']:s.get('org') for s in sources};url={s['id']:s.get('url') for s in sources}
    definitions=[
      dict(단계코드='POLICY',단계명='정책 수립',적용사업유형='정책·프로그램 전체',시작판정근거='공식 정책 원문과 발표일 확보',완료판정근거='정책 발표/고시 보관. 개별사업 완료와 다름',책임기관='관계부처·국토교통부 등; 개별 담당부서 미확인',공고채널='정책브리핑·부처 보도자료',원천문서유형='HTML/PDF/HWPX',수집방식='원문 파싱·목표/기간 정규화',비고='목표135만호를 각 사업별로 반복 합산하지 않음'),
      dict(단계코드='BUSINESS',단계명='사업화',적용사업유형='공공택지·정비·임대 등 유형별 정의 상이',시작판정근거='사업 명부·지구/블록·시행자·사업계획 확보',완료판정근거='선정/지정/사업화 상태의 실제 근거 또는 별도 표시한 담당자 추정',책임기관='시행기관·LH/SH/GH/iH·지자체·조합 등',공고채널='기관 사업현황·공모·고시',원천문서유형='CSV/XLSX/HTML/PDF/HWPX',수집방식='원천목록 파싱·기관별 필드 정규화',비고='사업 규모와 해당 회차 모집량을 구분'),
      dict(단계코드='APPROVAL',단계명='인허가',적용사업유형='주택사업계획승인·건축허가·정비 인가 등',시작판정근거='승인·인가·허가 등 사건 종류를 갖는 공시/원천',완료판정근거='사건일/상태/범위 확인. 없으면 추정으로 별도 표시',책임기관='관할 승인권자·지자체·국토부 등',공고채널='토지이음·지자체 고시·건축HUB·공식 월별통계',원천문서유형='API/XML/JSON/PDF/HWPX',수집방식='ID·사건종류·날짜·호수 파싱 및 개정 이력',비고='변경 인허가는 신규 추가공급과 다름. 통계 총계와 정책사업 명세를 분리'),
      dict(단계코드='CONSTRUCTION',단계명='착공',적용사업유형='건설형 주택 등',시작판정근거='착공 신고·실제 착공일·명시된 착공 상태/공정 근거',완료판정근거='착공 도달 사건. 공사계약일·입찰일과 구별',책임기관='시행/시공자·신고 수리기관·공공 발주기관',공고채널='건축HUB·LH 공사현황·정비사업·기관 공시',원천문서유형='API/CSV/DOCX/PDF/HTML',수집방식='블록/현장·실제일·관측상태 정규화',비고='이전 연도 착공도 정책 누적에 포함. 연간 정부 통계와 범위 상이'),
      dict(단계코드='COMPLETION',단계명='준공',적용사업유형='건설형 주택 등',시작판정근거='준공/사용승인 계획 또는 실제 사건',완료판정근거='실제 준공·사용승인 근거 확인',책임기관='시행자·검사/사용승인권자',공고채널='건축물대장·기관 준공계획/공시',원천문서유형='API/파일/공고',수집방식='착공과 별도 milestone; 계획·실제 분리',비고='공정률100%만으로 실제 입주를 확정하지 않음'),
      dict(단계코드='SUPPLY',단계명='공급·모집',적용사업유형='분양·임대·매입 등',시작판정근거='분양/모집 계획·공고',완료판정근거='실제 공고 및 금회 모집 범위 확인; 계약·입주와 구분',책임기관='공급기관·사업주체·청약/임대 공고기관',공고채널='청약홈·LH청약·SH/GH/iH 공고',원천문서유형='API/HTML/PDF/첨부',수집방식='공고ID·회차·금회/예비/잔여·정정 이력',비고='공급이 준공보다 앞설 수 있음. 국민/영구 임대는 별도 범위면 합산하되 중복 공고 제외'),
      dict(단계코드='MOVE_IN',단계명='입주',적용사업유형='주택공급 및 임대 유형별 상이',시작판정근거='입주 예정·개시 안내',완료판정근거='실제 입주 개시/현황 근거. 예정은 실적으로 변환하지 않음',책임기관='시행/공급기관·주택 관리기관',공고채널='입주 안내·모집공고·기관 일정',원천문서유형='파일/공고/HTML/PDF',수집방식='월·분기·일자 정밀도와 계획/실제 분리',비고='전체 실제 입주호수는 미측정 상태를 유지')]
    csvout('01_process_sources/supply_process.csv',definitions)
    workbook('01_process_sources/01_supply_process_and_sources.xlsx',[('단계_내부처리정의',definitions),('실제수집출처',registry),('집계단위',[dict(unit=k,definition=v) for k,v in {'사업 원천 기록':'기관별로 저장한 project 기록; 같은 현장 여러 원천 가능','관리사업':'원천별 중복과 범위를 처리한 집계 단위','호':'주택 물량; 공고건수/파일수/행정과제와 구분','공고 기록':'원천ID 단위; 기관 간 전수 중복제거 고유공고 수 아님','원문 해시':'보관한 응답바이트의 동일성; 사업 동일성 판정이 아님'}.items()])])
    md('01_process_sources/definitions.md','''
# 정책–프로그램–사업과 단계의 의미

이것은 the주택의 **내부 관리 모델**이며 모든 공급유형이 같은 순서로 진행되는 단일 법정 절차라고 주장하지 않는다.

- 정책: 공식 발표 원문과 범위·기간·목표를 가진 상위 관리단위.
- 프로그램: 유사 사업군을 정책에 연결하는 중간 단위. 자동 제안과 담당자 검토 상태를 보유한다.
- 사업 원천: 기관별 고유ID·제목·지구·블록·유형을 가진 수집 기록.
- 관리사업: 같은 현장·겹치는 원천을 정리한 정책 집계 단위. 별도 공급유형·부분 범위를 구분한다.

현재 주 매핑 테이블은 한 사업 원천당 선택 프로그램1개, 프로그램당 선택 정책1개를 저장한다. 여러 근거 문서와 다수 원천이 한 관리사업에 연결될 수 있지만 **주 정책·프로그램 귀속이 임의의 다대다로 저장되는 구조는 아니다.** 후속/관련 정책·중복불가 등 별도 관계 기록을 주 집계 귀속과 구분한다.

관측일은 사건이 그날 발생했음을 뜻하지 않을 수 있다. 계획의 월·분기·반기 말 환산일과 정확한 실제일도 구분한다. `baseline`은 담당자 채택 기준선일 수 있어 항상 최초 공표계획이라고 부르지 않는다.

선행단계 근거가 없을 때 후속 단계에서 선행 도달을 추정하는 기능이 있으나, 추정과 원천 실측을 구분해 저장한다. 공급은 준공 전에 공고될 수 있고 매입·정비·민간 유형은 행정절차가 다르므로 화면 흐름을 법적 선후관계로 단정하지 않는다.
''')
    batch=rows(OUT/'03_pipeline/pipeline_runs.csv');actions=rows(OUT/'03_pipeline/audit_action_counts.csv');ratio=round(capture['batch_status_counts']['SUCCESS']/capture['collection_batches']*100,3)
    progress=[dict(metric='성공 수집경로(누적1회 이상)',numerator=capture['sources_with_any_success'],denominator=capture['source_count'],unit='등록 경로',interpretation='등록 경로 중 성공 이력; 공개자료 전체 커버리지 아님'),dict(metric='성공 실행상태 비율',numerator=capture['batch_status_counts']['SUCCESS'],denominator=capture['collection_batches'],unit='수집 실행',interpretation='재시도·같은 원천 반복 포함; 문서 파싱 성공률 아님'),dict(metric='9·7 정책 연결 원천 비율',numerator=capture['linked_source_records'],denominator=capture['project_source_records'],unit='저장 사업 원천 기록',interpretation='전체 공개 사업의 정책 포함률 아님. REVIEW 포함 여부를 상태별 표와 함께 확인'),dict(metric='4분기 계획 사업명 대응',numerator=32,denominator=32,unit='공식 발표 행',interpretation='주택 수량/실제 공급 완료 일치와 구분'),dict(metric='4분기 원수량 동등',numerator=26,denominator=31,unit='중복 정리한 비교단위',interpretation='전체 데이터셋 정확도와 구분'),dict(metric='문서 정규화 성공률',numerator=None,denominator=None,unit='고유 원문',interpretation='미측정: 동일한 문서ID로 모든 추출 단계를 잇는 전수 로그 분모 미구축')]
    csvout('10_story_assets/measured_progress_bars.csv',progress)
    md('03_pipeline/pipeline_overview.md',f'''
# 실제 수집·전처리·연결 파이프라인

수집은 원문 확보·파싱·항목 정규화까지다. 정책/프로그램/사업 귀속 검토는 다음 연결 단계다.

```mermaid
flowchart LR
 A[등록 출처·목록/파일 탐색] --> B[수집 큐·LATEST/BACKFILL]
 B --> C[원문 HTTP 응답 보관·SHA-256]
 C --> D[API·HTML·파일별 파싱]
 D --> E[원천ID·필드 정규화·판본 저장]
 E --> F[현장·블록·유형 후보 연결]
 F --> G[담당자 확인·보류·반려]
 G --> H[정책 누적·계획 집계]
 H --> I[API·지도·보고서·CSV]
```

## 모듈별 입력·출력과 실패 처리

|모듈|실제 입력과 출력|실패·한계|
|---|---|---|
|예약/작업 큐|source_id→COLLECT→LINK→SNAPSHOT|실패 재시도·임대 만료 복구. 원문별 모든 처리단계와 일대일인 이벤트 로그는 아님|
|Fetcher|기관 허용URL→보관 응답·SHA·수집 manifest|429/503·일시 통신 오류 재시도·Retry-After. 오류를0건으로 저장하지 않음|
|API 파서|XML/JSON 응답→원천 레코드|페이지·지역 범위 확인. 반환 행수와 모집단 완결은 별개|
|HTML|BeautifulSoup 본문/링크→공고·첨부 참조|게시판 구조 변경은 파서 점검 필요|
|PDF|pypdf 텍스트/layout→자료별 정규식·표 규칙|**스캔 PDF 자동OCR은 미구현.** OCR_REQUIRED·보완 대상으로 남김|
|HWPX/DOCX|ZIP 내부 XML 문단·표→구조화 행|구형HWP는 보관/별도 보완; 모든 서식을 읽는 범용 엔진 아님|
|CSV/공개 엑셀|공개 파일/배포 포맷별 규칙→사업/계획|파일 링크·판본 변경 감지와 신규 데이터셋 발견은 다름|
|판본 저장|원천 식별키+정규화 payload hash→new/changed/unchanged|URL동일성·콘텐츠동일성·사업동일성은 각각 별개|
|연결/집계|블록·지역·유형·저장 근거→후보/확인/유보|공식 출처여도 제목만 같으면 무조건 자동 확정하지 않음|
|지역 표시자료|내부DB→지역·시군구·연도별 사전 계산 JSON|별도 작업자5분 변경 확인, 마지막 완료본 제공. 원천 API 수집 주기와 구분|

## 실제 로그에서 산출할 수 있는 수치

수집실행 {capture['collection_batches']:,}건 중 SUCCESS {capture['batch_status_counts']['SUCCESS']:,}건, FAILED {capture['batch_status_counts']['FAILED']:,}건, INTERRUPTED {capture['batch_status_counts']['INTERRUPTED']:,}건. 이 정의의 성공 비율은 {ratio}%다. 같은 자료의 반복 실행·재시도가 들어간 **실행 상태 비율**이며 고유 문서 성공률이나 정책 사업 수집률이 아니다.

- `pipeline_runs.csv`: 전체 실행 이력. 입력/실패/검토 건수가 원 기록에 없으면 공란(미측정).
- `processing_events.jsonl`: 원천 후보 발견 대량 이벤트와 독립스냅샷 등을 제외한 최근 {capture['audit_exported']:,}개 감사 이벤트 표본. 감사 전체 {capture['audit_total']:,}개 중 추출 범위를 `audit_action_counts.csv`와 함께 확인한다.
- `collection_jobs.csv`: 큐 작업·단계·재시도 상태. 수집 배치와 작업 수를 합산하지 않는다.
- `status_history.csv`: 일부 curated kind의 revision_no>1 판본 메타·선택 필드. 변경 사건수와 사업수를 구분한다.
- 원문마다 다운로드→추출→정규화→후보→승인을 같은 document_id로 잇는 **전수 단계별 처리율은 미측정**이다. 기존 batch·revision·audit·근거 링크의 결합으로 재현 가능한 부분을 제공한다.

## 변경 감지의 범위

보관 응답 해시가 반복되면 같은 바이트를 새 파일로 중복 생성하지 않는다. 정규화된 레코드 내용이 같으면 판본을 만들지 않고, 바뀌면 이전ID·superseded_at·새판본을 남긴다. 동일URL에서 다른 응답 해시가 나온 {capture['same_url_changed_response_candidates']:,}개 후보에는 API 응답 변화도 포함되므로 모두 정정공고/파일교체라고 표현하지 않는다.

등록된 문서를 재확인하는 수집기와 게시판 신규 목록을 탐색하는 수집기는 출처 대장에서 구분한다. 파서 규칙 밖의 이미지·변경서식은 수동 확인 대상이며 생성형 AI/OCR 모델이 모두 자동 처리한다고 주장하지 않는다.
''')
    md('04_entity_linking/relation_rules.md','''
# 연결 키와 담당자 확인의 실제 의미

## 데이터 키

`policy_id`·`program_id`·사업 원천ID를 내부 키로 유지하고 원천 기관ID/공고ID/지구/블록/지역/유형을 별도 보유한다. ProjectIdentity와 IdentityAlias는 동일 현장 대응, Mapping/ProgramMapping은 정책 귀속, stage_evidence/evidence_bridge/collection_link는 단계·문서 근거다. 파일 SHA-256으로 현장 동일성을 확정하지 않는다.

## 자동과 담당자 확인

일반 `MAPPING_AUTO`는 신규 원천을 **REVIEW 후보**로 제안하는 경로를 포함한다. 저장된 상태와 명시적 검토 이력, 출처·지구·블록·물량 범위를 확인한 규칙이 확정한 경우를 구분해야 한다. “공식이면 전부 자동 확정”이라는 설명은 부정확하다.

비공식 근거는 기존 사용자 지시에 따라 **담당자 확인 - 비공식 자료: 출처**로 관리한 실제 DB 기록이 있다. 이는 원문을 공식 자료로 등급 변경한 것이 아니다. 일부 확인은 사용자 위임을 받은 에이전트가 담당자 역할을 대행해 등록한 것이며 **행정기관 실무자의 실제 인적 결재 증명으로 제출할 수 없다.** actor 역할 표기만 있는 과거 기록 역시 실제 사람의 결재 여부를 추정하지 않는다.

`match_review.csv`는 추출한 감사 표본에서 확인 관련 이벤트를 고른 기록이다. 모든 후보의 최초 제출시각이 별도 저장된 것은 아니며, 제출시각 공란을 결정시각으로 복제하지 않는다. 현재 상태의 건수와 이력상 승인 이벤트 수는 같은 분모가 아니다.

## 합산 원칙

- 동일 현장 공사·공정률·공급계획의 전체량은 서로 더하지 않는다.
- 같은 블록의 국민임대/영구임대 등 서로 겹치지 않는 유형은 구성 확인 후 합산 가능하다.
- 전체 지구·공구·일부 블록과 모집회차·금회/예비자 물량을 구별한다.
- 같은 사업의 선행·후속 단계 도달량은 단계 간 중복이므로 합산하지 않는다.
- 후속 단계에서 선행 도달을 추정한 경우 `inferred`·근거·원천과 구분한다.
- 수량을 모르는 사업은0호가 아니며 관계를 연결했다고 정책 총량을 반드시 가산하는 것도 아니다.

대표 연결 난점: 연남8호와 연희96호는 계약 제목만으로 같은 현장이라고 볼 수 없었다. 부평 A는 단독 문자 블록 파싱 차이 때문에 두 원천을 다른 현장으로 볼 수 있었다. 상세 근거·전후 숫자는 `05_case_studies`에 제공한다.
''')
    status=Counter(p.get('mapping',{}).get('status') for p in projects);policy_status=Counter(p['mapping']['status'] for p in projects if p.get('mapping',{}).get('policy_id')=='POL_2025_0907')
    csvout('04_entity_linking/current_status_counts.csv',[dict(scope='전체 원천사업',status=k,count=v,unit='현재 사업 원천 기록') for k,v in status.items()]+[dict(scope='9·7 귀속 표기',status=k,count=v,unit='현재 사업 원천 기록') for k,v in policy_status.items()]+[dict(scope='자동 확정·실제 사람 결재 분리',status='미측정',count=None,unit='행위 주체를 전수 구분한 독립 명부 없음')])
    return definitions,registry,progress

def cases(projects,quantity,starts):
    basedir=OUT/'05_case_studies';basedir.mkdir(exist_ok=True)
    case_rows=[]
    configs=[
    ('CASE-001','연남동8호와 연희96호를 별도 사업으로 재식별','동일사업매칭','96호 일부라는 과거 가설','연남8호 독립사업; 연희96호는 별도','8호는 기존 원장 유지, 실제 착공/준공 추가0건','docs/HOME_YEONNAM_8_IDENTITY_REVIEW_20261020.md',['docs/analysis/yeonnam_sh_identity_20261020/verified_rows.json','docs/analysis/yeonnam_sh_identity_20261020/source_manifest.json','docs/analysis/yeonnam_sh_identity_20261020/comparison_verification.json'],'공식PDF79쪽 텍스트 추출·별도 행/규모/시행자 대조와 페이지 시각 확인. 이미지OCR 성과가 아니다.'),
    ('CASE-002','인천부평 A의 공사·공정률·계획 중복 제외','동일현장중복','전체지도 목표700호','전체지도 목표350호','전체지도−350호;9·7 관리350호는 유지','docs/HOME_BUPYEONG_DEDUP_20261020.md',['docs/analysis/bupyeong_dedup_20261020/before_after.json','docs/analysis/bupyeong_dedup_20261020/deployment.json'],'숫자가 없는 A단독 블록을 지구·지역과 함께 동일키로 정규화. 다른 A/B/A1은 분리.'),
    ('CASE-003','검단 AA19 국민·영구임대를 유형별로 검산','유형별합산','발표에서는 임대 약1,300호 묶음','국민991+영구332=1,323호','단일 지구와 두 구성유형을 중복 합산하지 않음','docs/HOME_GPT_REPORT_MASTER_20261020.md',[],'유형·블록·공고범위로 겹치지 않는 부분량을 구분. 사업 원천명은 current_source_projects.csv 참조.'),
    ('CASE-004','부천대장 A9의 승인 판본568→574','인허가변경','과거 승인568호','변경승인574호','발표 약600호와 규모 대응. 현재 누적에는 기존 다른 블록도 포함','docs/HOME_THREE_NEWTOWN_FINAL_COMPARISON_20261020.md',['docs/analysis/three_newtown_final_20261020/final_comparison.csv','docs/analysis/three_newtown_final_20261020/source_manifest.json'],'변경 승인과 착공 확인 자료의 역할을 분리. 계약일을 실제 착공일로 바꾸지 않음.'),
    ('CASE-005','고덕강일3 추정 공급시점과 실제 본청약 구분','계획실적보정','공급관리계획2026-08-31·일부 잔여량215호','본청약2026-10-16·전체1,305호','동일1,305호 규모의 실적 사건/범위를 보정, 추가1,305호 중복 가산 아님','docs/HOME_EXTERNAL_STAGE_EVIDENCE_APPLIED_20261020.md',['docs/analysis/reliability_update_20261022/official-32.csv'],'담당자 확인 - 비공식 자료: 뉴시스 등. 전체량1305와 금회/잔여량215를 같은 범위로 오인하지 않음.'),
    ('CASE-006','남양주왕숙2 공모·계약의19호 상충 보존','미해결판본상충','공모·확인 자료A6 510+A7 553=1,063호','계약공시A6 519+A7 563=1,082호도 존재','상충19호 보존; 두 값을 더하지 않음','docs/HOME_THREE_NEWTOWN_FINAL_COMPARISON_20261020.md',['docs/analysis/three_newtown_final_20261020/final_comparison.csv','docs/analysis/three_newtown_final_20261020/source_manifest.json'],'기존 A1/A3 1,498호와 다른 블록이다. 상충을 임의 반올림/가짜 정정으로 해소하지 않음.')]
    for cid,title,typ,before,after,impact,report,files,rule in configs:
        dest=basedir/cid;dest.mkdir(exist_ok=True)
        pattern={'CASE-001':lambda p:'연남' in p['name'] or '연희' in p['name'],'CASE-002':lambda p:'부평' in p['name'],'CASE-003':lambda p:'검단' in p['name'] and 'AA19' in p['name'].upper(),'CASE-004':lambda p:'부천대장' in p['name'] and 'A9' in p['name'].upper(),'CASE-005':lambda p:'고덕강일' in p['name'] and ('3' in p['name']),'CASE-006':lambda p:'왕숙2' in p['name'] and any(t in p['name'].upper() for t in ['A6','A7'])}[cid]
        ps=[p for p in projects if pattern(p)];csvout(f'05_case_studies/{cid}/current_source_projects.csv',[{k:p.get(k) for k in ('id','name','source_id','revision_no','district','block','tenure','planned_units','source_date','base_date','evidence_url','start_actual_date','recorded_at')} for p in ps])
        refs=[]
        for path in files:
            src=ROOT/path
            if src.exists():
                target=f'05_case_studies/{cid}/{src.name}'
                if src.suffix=='.json':dump(target,json.loads(src.read_text()))
                elif src.suffix=='.csv':csvout(target,rows(src))
                refs.append(src.name)
        if cid=='CASE-001':
            shutil.copyfile(A/'yeonnam_sh_identity_20261020/page-079.png',dest/'page-079.png');copy_text('docs/analysis/yeonnam_sh_identity_20261020/page-079.txt',f'05_case_studies/{cid}/page-079.txt');refs+=['page-079.png','page-079.txt']
        if cid=='CASE-003':
            verified=[r for r in rows(A/'reliability_update_20261022/official_185_comparison.csv') if r['지표ID']=='R063'];csvout(f'05_case_studies/{cid}/official_comparison.csv',verified);refs.append('official_comparison.csv')
        sourceurls=sorted({p.get('evidence_url') for p in ps if p.get('evidence_url')})
        # Include scoped confirmation entries already preserved in the 32-row ledger.
        if cid=='CASE-005':csvout(f'05_case_studies/{cid}/confirmation_ledger.csv',[r for r in quantity if '고덕강일' in r['official_name']]);refs.append('confirmation_ledger.csv')
        row=dict(case_id=cid,case_title=title,case_type=typ,source_ids=sorted({p.get('source_id') for p in ps}),source_urls=sourceurls,internal_project_ids=[p['id'] for p in ps],before_value=before,after_value=after,normalization_or_matching_rule=rule,approval_or_verification='담당자 대행 확인/공식 근거 대조; 실사용 기록과 실제 인적 결재 증명은 구분',impact_on_aggregate=impact,implement_status='실제 원장·검증기록 보유',can_show_publicly='개인정보 제거한 공개원천 메타·선정 페이지',origin_report=report,evidence_files=refs)
        dump(f'05_case_studies/{cid}/case.json',row);case_rows.append(row)
        md(f'05_case_studies/{cid}/README.md',f'''# {cid} · {title}

1. **원본:** 아래 원문URL·출처ID·판본과 `current_source_projects.csv`를 확인한다.
2. **문제 구간:** {before}.
3. **추출/과거값:** {before}.
4. **정제·확인 규칙:** {rule}
5. **결과:** {after}.
6. **서비스 영향:** {impact}.

이전값은 해당 검증 당시 값이며 현재 원천의 숫자를 임의로 과거값처럼 복제하지 않았다. 실제 법정 사업 분할/합병 사례와 데이터 중복 통합/유형별 부분합은 구분한다. 근거를 공개 범위로 추출한 사례이며 별도 행정기관 인적 승인 증명은 아니다.

원문URL:
'''+ '\n'.join('- '+u for u in sourceurls)+'\n\n동봉 근거:\n'+ '\n'.join('- ['+f+']('+f+')' for f in refs))
    csvout('05_case_studies/case_index.csv',case_rows)
    md('05_case_studies/없는_사례와_표현한계.md','''
# 입증하지 못한 사례

- 이미지뿐인PDF를 OCR/비전모델로 전량 자동 정형화한 성공사례: **미구현/입증불가**. 연남PDF 사례는 텍스트 추출과 원문 페이지 확인이다.
- 실제 법정사업1건이 여러 법정사업으로 분할·합병된 전후 행정명부 전수사례: **이번 제출에서 미확보**. 부평은 원천중복 통합, 검단은 공급유형 분해·합산 사례다.
- 같은 다운로드URL의 파일이 바뀐 것을 공식 정정공고라고 개별 확정한 전수사례: 미확보. 같은URL 응답해시 변경 후보 목록은 제공하되 API 응답 변화도 포함한다.
- 사람의 수작업 시간을 시간 단위로 측정한 기록: 미측정. 담당자 확인 기록 수를 인력투입시간으로 환산하지 않는다.
''')
    return case_rows

def quality_and_snapshots(capture,projects,model,summary,ledger,detailed,quantity):
    initial=old('docs/HOME_REQUIREMENTS_AUDIT_20261012.json')['snapshot']['projects']
    snap21=old('docs/analysis/policy907-full-verification-20261021/snapshot.json');snap22=old('docs/analysis/reliability_update_20261022/current-validation.json')
    snapshots=[dict(snapshot_id='REQ_20261012',as_of_date='2026-10-12',policy_count=None,program_count=None,project_count=initial['count'],notice_count=None,file_count=None,parsed_record_count=None,linked_project_count=None,unlinked_count=None,household_total=None,management_project_count=None,scope_description='당시 요구사항 점검의 전체 사업 원천 기록',filters='전체 수집 프로젝트;9·7 전수만 아님',code_version='해당 보관 JSON의 버전 참조',evidence_ref='references/requirements_20261012_snapshot.json'),
      dict(snapshot_id='CONSISTENT_20261021',as_of_date=snap21['snapshot_at'],policy_count=None,program_count=None,project_count=snap21['projects'],notice_count=None,file_count=None,parsed_record_count=None,linked_project_count=snap21['linked_source_projects'],unlinked_count=None,household_total=None,management_project_count=snap21['management_cohorts'],scope_description='읽기전용 전체 원천사업 및9·7 관리범위',filters='POL_2025_0907;당시 집계 규칙',code_version='기존 고정 추출본;현재 코드로 덮어쓰지 않음',evidence_ref='references/snapshot_20261021.json'),
      dict(snapshot_id='R02_REVIEW_20261022',as_of_date=snap22['assessed_at'],policy_count=None,program_count=None,project_count=None,notice_count=None,file_count=None,parsed_record_count=None,linked_project_count=snap22['contributing_source_ids'],unlinked_count=None,household_total=None,management_project_count=snap22['portfolio_cohorts'],scope_description='2차 성적표 재평가의 기여 원천221·관리177',filters='기여 원천 수;일반 연결원천232와 분모 다름',code_version='공식점수/내부준비도 평가',evidence_ref='references/current_validation_20261022.json'),
      dict(snapshot_id='PRESENTATION_EXTRACT_20261022',as_of_date=capture['extracted_at'],policy_count=capture['policy_count'],program_count=capture['program_count'],project_count=capture['project_source_records'],notice_count=capture['notice_records'],file_count=capture['unique_response_hashes'],parsed_record_count=None,current_revision_record_count=capture['current_records'],linked_project_count=capture['linked_source_records'],unlinked_count=None,household_total=sum(c['units'] for c in model['portfolio_cohorts']),management_project_count=capture['management_cohorts'],scope_description='현재 원천452·정책연결232·기여221·관리177;file_count는API포함 보관응답 해시',filters='전체 출처 및9·7 관리집계;필드별 범위 구분',code_version='code_version.json의미커밋작업트리해시',evidence_ref='00_data/capture_summary.json')]
    csvout('06_data_quality/dataset_snapshots.csv',snapshots)
    dump('references/requirements_20261012_snapshot.json',initial);dump('references/snapshot_20261021.json',snap21);dump('references/current_validation_20261022.json',snap22)
    from collections import Counter
    q=[]
    def check(id,name,n,passed,basis):q.append(dict(check_id=id,snapshot_id='PRESENTATION_EXTRACT_20261022',check_name=name,input_scope='명시된 추출 범위',expected_rule=basis,records_checked=n,records_passed=passed,records_failed=n-passed,issue_id=None,resolution_status='PASS' if n==passed else 'REVIEW',evidence_ref='09_map_dashboard/project_stage_counting_ledger.csv'))
    cohorts=model['portfolio_cohorts'];contributing=[x for c in cohorts for x in c['source_ids']]
    check('Q01','기여 원천ID를 관리사업 간 중복 사용하지 않음',len(contributing),len(set(contributing)),'source_ids 고유성')
    check('Q02','관리사업ID 고유성',len(cohorts),len({c['id'] for c in cohorts}),'관리단위 중복 금지')
    check('Q03','현재량은 관리사업 상한 이내',len(ledger),sum(r['units'] is None or 0<=r['units']<=r['capacity'] for r in ledger),'null은미확보;도달량<=관리상한')
    check('Q04','연말 누적예상>=현재 누적',len(ledger),sum(r['units'] is None or (r['year_end_units'] is not None and r['year_end_units']>=r['units']) for r in ledger),'현재와동일범위·연말까지 저장계획 합집합')
    check('Q05','공식32행의 현재 관리ID 대응',32,sum(r['home_cohort_id'] in {c['id'] for c in cohorts} for r in quantity),'사업명 대응이며 실제공급 완료검증 아님')
    check('Q06','185개 판정 상태 미공란',185,sum(bool(r['판정']) for r in detailed),'판정 상태 기록;일치수 아님')
    check('Q07','착공 관리합계 재현',1,int(next(r['current_units'] for r in summary if r['stage']=='CONSTRUCTION')==82968),'177관리사업의누적도달량합계=82968')
    check('Q08','공식32행 단순행수와 수량비교단위 구분',1,int(len({r['home_cohort_id'] for r in quantity})==31),'A1/A3 한묶음으로31비교단위')
    csvout('06_data_quality/quality_checks.csv',q)
    w=__import__('openpyxl').load_workbook(A/'housing_supply_97_official_verification_20261021.xlsx',read_only=True,data_only=True)
    sh=w['교차검증_쟁점'];it=iter(sh.values);headers=next(it);issues=[{str(k):v for k,v in zip(headers,line) if k is not None} for line in it if any(v is not None for v in line)];csvout('06_data_quality/issue_log.csv',issues)
    copy_text('docs/analysis/THE_JUTAEK_INTEGRITY_R02_20261022.md','06_data_quality/scorecard_latest.md')
    md('06_data_quality/snapshot_notes.md','''
# 스냅샷과 품질검증의 범위

처음 크롤링·처음 정형화·처음 관계매핑 시점마다 모든 지표를 고정한 동일형식 스냅샷은 미확보다. 존재하는 날짜별 점검자료와 읽기전용 추출본만 제공한다. 공란을 최신값으로 채워 과거 실적을 만들지 않았다.

source 사업수와관리사업수,정책 연결원천과실제집계기여원천의분모가다르다. 2026-10-12 292와최신452는전체원천기록비교이며177관리사업증가율과같지않다. metric_snapshot_register는별도독립집계장부의고정기록으로화면정책누적과다를수있다.

quality_checks는이번패키지에실행한좁은범위의집계·추적검사다. 데이터전수진실성검증이아니다. 추가로보존된최초/1차/2차성적표3개는integrity_rounds.json에원래점수와검토시각을유지한다. 기존2차63.2/100,공식비교12.95/20은다른척도이며정부인증정확도점수가아니다.
''')
    return snapshots,q,issues

def starmap(capture,sources,projects,programs):
    nodes={};edges=[]
    def node(id,typ,label,**kwargs):
        nodes.setdefault(id,dict(node_id=id,node_type=typ,label=label,source_id=None,document_id=None,project_id=None,file_format=None,step=None,record_count=1,observed_at=capture['extracted_at'])|kwargs)
    def edge(fr,to,rel,method='보관된 원천 필드/관계',status='METADATA',evidence=None):
        if fr in nodes and to in nodes:edges.append(dict(edge_id='EDGE_'+hashlib.sha256((fr+'|'+to+'|'+rel).encode()).hexdigest()[:20],from_node_id=fr,to_node_id=to,relationship=rel,source_evidence_id=evidence,connection_method=method,approval_status=status))
    for s in sources:
        oid='ORG_'+hashlib.sha256(str(s.get('org')).encode()).hexdigest()[:16];host=urlsplit(s.get('url','')).hostname or '미기재';site='SITE_'+host
        node(oid,'기관 표기',s.get('org'));node(site,'사이트',host);node('SRC_'+s['id'],'등록 수집경로',s['name'],source_id=s['id']);edge(oid,'SRC_'+s['id'],'경로 등록 기관');edge(site,'SRC_'+s['id'],'대표 접근 도메인')
    for p in read('00_data/policies.json'):node(p['policy_id'],'정책',p['name'])
    for p in programs:node(p['id'],'프로그램',p['name']);edge(p.get('policy_id'),p['id'],'정책-프로그램 저장 귀속',status=p['status'])
    for p in projects:
        node(p['id'],'사업 원천',p['name'],project_id=p['id'],source_id=p.get('source_id'));edge('SRC_'+p.get('source_id',''),p['id'],'사업 원천 수집');edge(p.get('mapping',{}).get('program_id'),p['id'],'프로그램-사업 저장 귀속',status=p.get('mapping',{}).get('status'))
    docs=rows(OUT/'02_collection_starmap/document_inventory.csv');docids=set()
    # Complete curated document nodes; separate large notice-catalogue list from body metadata.
    for d in docs:
        if d['record_kind']=='notice_catalogue':continue
        node(d['document_id'],'문서/공고 기록',d['title'],source_id=d['source_id'],document_id=d['document_id']);docids.add(d['document_id']);edge('SRC_'+d['source_id'],d['document_id'],'원천 기록 수집')
    for r in rows(OUT/'04_entity_linking/relations.csv'):
        if r.get('review_status') not in ('REJECTED','RELEASED','UNMATCHED'):edge(r['parent_id'],r['child_id'],r['relation_type'],r.get('method'),r.get('review_status'),r.get('source_document_id'))
    selected=[];formats=Counter();observations=0;pdfs=[];changed_content=[]
    for r in rows(OUT/'02_collection_starmap/file_inventory.csv'):
        exts=json.loads(r['stored_extensions']);
        for ext in exts:formats[ext]+=1
        observations+=int(r['fetch_observations']);
        if len(selected)<1500 or any(ext in ('pdf','hwpx','hwp','xlsx','xls','docx','png','jpg') for ext in exts):
            if len(selected)<3000:selected.append(r)
    for r in selected:
        node(r['file_id'],'보관 응답파일 표본',r['sha256'][:14],file_format=r['stored_extensions'])
        for sid in json.loads(r['source_ids']):edge('SRC_'+sid,r['file_id'],'응답 파일 보관',evidence=r['sha256'])
    # Only link document→file where an explicit original hash agrees.
    bysha={r['sha256']:r['file_id'] for r in selected}
    for d in docs:
        if d['original_sha256'] in bysha:edge(d['document_id'],bysha[d['original_sha256']],'문서 original_sha256 일치',evidence=d['original_sha256'])
    unique={e['edge_id']:e for e in edges};csvout('02_collection_starmap/nodes.csv',nodes.values());csvout('02_collection_starmap/edges.csv',unique.values())
    stats=[dict(metric='현재 저장 레코드(원천·관계·관리 포함)',value=capture['current_records'],unit='기관별 current record',scope='API 대장·지역통계·사업·문서 등 전체. 주택 호수 아님'),dict(metric='등록 수집경로',value=capture['source_count'],unit='경로',scope='실제 레지스트리'),dict(metric='등록기관 표기 고유수',value=capture['source_organization_count'],unit='기관 문자열',scope='기관 조직코드로 정규화한 기관수와 다름'),dict(metric='대표 접근 도메인 고유수',value=capture['source_domain_count'],unit='hostname',scope='등록 출처 대표URL. 게시판 수 아님'),dict(metric='공고 기록',value=capture['notice_records'],unit='현재 notice/lh_notice/gh_notice 원천ID',scope='기관 간/정정회차간 전수 중복제거 고유공고수와 구분'),dict(metric='수집·검토 문서 근거',value=capture['document_records'],unit='document_evidence ID',scope='채택/미연결/검토 포함'),dict(metric='목록·문서·공고 inventory 전체',value=capture['document_inventory_records'],unit='정규화 기록',scope='notice_catalogue 등 목록 레코드도 포함'),dict(metric='첨부 참조',value=capture['attachment_references'],unit='문서 attachments 참조행',scope='다운로드 성공파일/고유첨부수와 구분'),dict(metric='응답 콘텐츠 고유해시',value=capture['unique_response_hashes'],unit='SHA-256',scope='API JSON/XML·HTML 포함, PDF문서 개수가 아님'),dict(metric='manifest 응답 관측',value=capture['manifest_observations'],unit='다운로드 관측',scope='같은 바이트 반복 포함'),dict(metric='동일 응답 반복 관측',value=capture['repeat_response_observations'],unit='추가 관측',scope='동일 콘텐츠 재확인. 중복사업 제거 건수 아님'),dict(metric='저장 확장자 종류',value=len(formats),unit='확장자',scope='MIME/스캔 PDF 식별 수와 다름'),dict(metric='고유 게시판 수',value=None,unit='게시판',scope='전수 식별자 미구축; source 명칭/도메인수로 대체하지 않음'),dict(metric='스캔PDF/OCR 완료율',value=None,unit='PDF',scope='미측정/자동 OCR 미구현'),dict(metric='스타맵 노드',value=len(nodes),unit='선택 노드',scope='등록경로·curated 문서/사업 전수 및파일 표본'),dict(metric='스타맵 연결선',value=len(unique),unit='보관 관계',scope='양끝 노드가 있는 명시적 관계. 후보 상태와확정 상태 구분')]+[dict(metric='보관확장자 '+fmt,value=n,unit='고유 해시',scope='하나의 해시에 여러 확장자가 있으면 중복될 수 있음') for fmt,n in formats.most_common()]
    csvout('02_collection_starmap/source_stats.csv',stats)
    md('02_collection_starmap/README.md',f'''
# 스타맵용 실데이터와 해석

nodes.csv는{len(nodes):,}개노드,edges.csv는{len(unique):,}개관계이다. 수집경로·정책·프로그램·사업 원천 및현재 curated 문서/공고를포함하고,대량API응답파일은{len(selected):,}개표본만넣었다. 원문파일메타전수는file_inventory.csv에있다.

연결선은저장된필드·관계에근거하며후보/확정상태를보존한다. 파일과문서는original_sha256이일치할때만직접연결했다. 문서의인허가값이파일어느페이지어느셀에서추출됐는지전수좌표가없는경우그관계를새로만들지않았다.

보관확장자json/xml/html등은HTTP응답형식이다. 이를모두비정형공고파일이라고부르지않는다. 해시는인증값제거후보관응답의해시일수있으며,공개PDF원본해시가별도확보된경우와구분한다.

게시판고유수·전수PDF스캔분류·원문별전체처리전환율은미측정이다. 수동파싱보완의전체업무시간도숫자로만들지않는다.
''')
    return stats

def system_and_dictionary():
    import ast
    from backend.db import Base
    fields=[]
    for t in Base.metadata.sorted_tables:
        for c in t.columns:fields.append(dict(table=t.name,column=c.name,type=str(c.type),nullable=c.nullable,primary_key=c.primary_key,foreign_key='|'.join(str(k.target_fullname) for k in c.foreign_keys),note='payload JSON의 가변 필드는 별도 원천/단계 CSV 참조' if c.name=='payload' else ''))
    csvout('12_technical/data_dictionary.csv',fields)
    paths=[];tree=ast.parse((ROOT/'backend/app.py').read_text())
    for f in ast.walk(tree):
        if not isinstance(f,(ast.FunctionDef,ast.AsyncFunctionDef)):continue
        for dec in f.decorator_list:
            if isinstance(dec,ast.Call) and isinstance(dec.func,ast.Attribute) and isinstance(dec.func.value,ast.Name) and dec.func.value.id=='app' and dec.args:
                p=dec.args[0]
                if isinstance(p,ast.BinOp) and isinstance(p.right,ast.Constant):path='/poc/home'+str(p.right.value)
                elif isinstance(p,ast.Constant):path=p.value
                else:continue
                paths.append(dict(method=dec.func.attr.upper(),path=path,handler=f.name,authentication='관리자 인증 검사' if any(isinstance(n,ast.Call) and isinstance(n.func,ast.Name) and n.func.id=='authenticated' for n in ast.walk(f)) else '공개 조회 또는 별도 인증경계',code_ref='backend/app.py:'+str(f.lineno)))
    csvout('12_technical/api_inventory.csv',paths)
    roles=[dict(actor='업무담당자',actual_flow='원천확인→후보/근거조회→연결/단계확인→보고',actual_screens='출처와 수집현황·사업별 수집연계·정책-프로그램-사업 연결·사업 상세·보고서',permission='저장/변경 API는 관리자 인증 검사. 이번 캡처는 비로그인 읽기모드',implementation='기능 구현;실제 기관 담당자 인적 결재 여부 별도',network='현재 PoC 같은 서비스/관리자 인증 경계. 실제 내부망 분리 운영 미확인'),dict(actor='정책담당자',actual_flow='정책/프로그램/사업 지도→목표/실적/연말계획→점검/시나리오',actual_screens='정책 추진현황·프로그램 추진현황·사업 추진현황·지역 여건',permission='읽기 조회·관리자 저장 기능은 구분',implementation='현재 PoC 구현',network='역할 전용 별도 인증체계3종으로 분리된 것은 아님'),dict(actor='국민',actual_flow='관심 지역·사업 조회→일정·원문 확인',actual_screens='공급 현황 지도·지역 및사업필터·사업 위치/주변정보',permission='공개 조회',implementation='현재 서비스의 국민 관점 활용. 별도 국민서비스 배포 검증은 아님',network='공식 확인 자료만의 외부 공개 승인망은 현재 PoC에서 전수 입증되지 않음')]
    workbook('08_service/08_roles_and_permissions.xlsx',[('3액터_실제범위',roles),('API경계',paths)])
    csvout('08_service/roles_and_permissions.csv',roles)
    md('08_service/screen_inventory.md','''
# 실제 화면·권한과 미구현 구분

캡처는 현재 PoC10의 실제 서비스 데이터로 수행한다. 관리자 계정을 새로 만들거나 승인·데이터 변경을 실행하지 않는다.

|관점|제공 화면|상태|
|---|---|---|
|업무담당자|사업별 수집연계·출처·연결 목록|실제 조회 화면, 비로그인 읽기모드|
|정책담당자|정책·프로그램·사업 추진현황, 보고서 미리보기|실제 화면|
|국민|지역 공급지도·사업 및 일정·주변정보|동일 PoC의 국민 관점 조회화면|

현재 구현은 관리용 API·수집 작업자·지역 가공 작업자·DB를 분리한 동일 PoC 서비스다. 최초 외부망 수집 설계안에서 내부망 업무플랫폼으로 이동했다는 **실제 배포 연혁/네트워크분리 증빙은 미확보**다. 세 액터가 서로 다른 목적을 갖는다는 개념과 실제3개 독립 서비스·내외망 분리 운영은 같은 주장이 아니다.

개인조건별 공고 자격 판정·개인 맞춤 추천·전국 모든 사업전수·모든 비공식 후보를 차단한 외부발행 승인망은 이번 코드/화면의구현성과로주장하지않는다. 기본 화면의 '전국' 관점은 이 PoC의 수도권3개시도 요약으로설명한다.

보고서 만들기는선택한정책/프로그램/사업에맞춰미리보기·브라우저PDF저장·MD를제공한다. 텍스트를새로생성한AI정책판단모형이나공식정부작성보고서가아니다.
''')
    md('12_technical/architecture.mmd','''
erDiagram
  POLICY ||--o{ PROGRAM_MAPPING : "선택 정책 귀속"
  PROGRAM_MAPPING ||--o{ MAPPING : "선택 프로그램 귀속"
  MAPPING ||--|| SOURCE_PROJECT : "원천 사업ID"
  PROJECT_IDENTITY ||--o{ IDENTITY_ALIAS : "동일 현장 검토"
  IDENTITY_ALIAS }o--|| SOURCE_PROJECT : "원천 대응"
  SOURCE_PROJECT ||--o{ REVISION : "내용 판본"
  SOURCE_PROJECT ||--o{ STAGE_EVIDENCE : "단계 확인"
  SOURCE_PROJECT ||--o{ EVIDENCE_BRIDGE : "문서 연결"
  SOURCE_PROJECT ||--o{ EXECUTION_SCHEDULE : "목표/계획"
  SOURCE_PROJECT ||--o{ AUDIT : "확인/변경 이력"
  BATCH ||--o{ REVISION : "수집 실행"
''')
    md('12_technical/implementation_and_limits.md','''
# 구현의 실제 범위

- Python/FastAPI/SQLAlchemy/PostgreSQL(PostGIS), 브라우저JavaScript·지도·3D라이브러리 사용. 라이브러리 버전은 requirements.txt와 코드해시 목록 참조.
- pypdf 텍스트 추출,BeautifulSoup HTML,ZIP/XML HWPX/DOCX,기관별API/XML/JSON 및파일 파서가있다. 통합OCR/비전모델로스캔PDF를전량자동처리한기능은없다.
- 원천별source_object_id와내부해시ID,이전판본·payload_hash·superseded_at,감사before/after를보관한다. 원문SHA는콘텐츠동일성,정책/사업ID는귀속과동일현장판단이므로서로대체하지않는다.
- 지연등급은도래목표대비현재측정부족률기준으로0초과~2%미만주의,2%이상지연이다. 정확한일정이없거나미측정이면판단보류로구분한다. 이것은모든현장병목원인을자동진단한성과가아니다.
- 병목/조치등록·일정이동시나리오및관리화면이있어도모든기관의실제행정처리이력이들어온것은아니다.
- 가구·주택재고의산술차이는실제공가·가용주택·확정주택수요부족량이아니다.
- 최초계획일·변경계획일·실제일이모든사업에완결돼있지않다. 담당자채택baseline을최초공표계획으로일괄치환하지않는다.

제도개선의구체적근거는연남8/연희96의공통키부재,부평A블록동일현장파싱,검단유형별부분합,대장568/574변경판본,왕숙1063/1082상충이다. 공고에공통사업ID·정책/프로그램귀속ID·단계/사건코드·호수범위·정정대상ID·원문URL/해시·공표일/사건일을함께주면이들문제의재발을줄일수있다.
''')
    md('08_service/architecture.html','''<!doctype html><html lang="ko"><meta charset="utf-8"><title>the주택 시스템 구성</title><style>
body{font:20px/1.6 sans-serif;color:#183c55;background:#f0f5f8;margin:0;padding:45px;max-width:1450px}h1{font-size:36px;margin:0}p{color:#526b7f}.flow{display:grid;grid-template-columns:repeat(4,1fr);gap:20px;margin:28px 0}.box{border:1px solid #c5d8e7;background:white;border-radius:14px;padding:24px;min-height:170px}.box b{display:block;color:#12598c;font-size:24px}.box small{font-size:16px}.roles{display:grid;grid-template-columns:repeat(3,1fr);gap:22px}.roles .box{background:#e3f1f3;border-color:#72a8b2;min-height:120px}.arrow{font-size:22px;text-align:center;color:#1c709a}.note{padding:20px;border-left:5px solid #b28a33;background:#fff8e8;font-size:17px}</style>
<h1>the주택 · 공개 원천에서 지도·보고서까지</h1><p>현재 PoC10 구현을 기준으로 한 구성도</p>
<div class="flow"><div class="box"><b>① 공개 원천</b>정부·지자체·공공기관<br>API·공고·보도·첨부파일<br><small>뉴스는 비공식 근거로 구분</small></div><div class="box"><b>② 수집 작업자</b>예약·재시도·범위 이어받기<br>다운로드 → 파싱·정규화<br><small>스캔PDF 자동OCR 미구현</small></div><div class="box"><b>③ 내부 DB</b>원천 판본·해시·감사이력<br>정책 → 프로그램 → 사업<br><small>후보·담당자 확인·추정 구분</small></div><div class="box"><b>④ 조회·보고</b>API·지도·단계·시계열<br>MD·PDF·CSV<br><small>지역 자료는 별도 작업자로 사전 계산</small></div></div>
<div class="arrow">각 관점에 맞는 화면 · 실제 근거로 원천까지 추적 ↓</div><div class="roles"><div class="box"><b>업무담당자 관점</b>원천·후보 확인 / 연결·변경<br>관리자 저장 인증</div><div class="box"><b>정책담당자 관점</b>목표·실적·계획 비교<br>지역·프로그램·사업 점검</div><div class="box"><b>국민 관점</b>지역 공급정보·일정·원문<br>사업 위치와 주변 여건</div></div>
<p class="note">실제3개 독립서비스·내부망/외부망 분리운영을 입증하는 구성도가 아닙니다. 현재 같은 PoC 안의 관점별 활용과 관리자 저장 경계를 표현했습니다. 에이전트 담당자 대행 기록을 행정기관 인적 결재로 표시하지 않습니다.</p></html>''')
    return roles,paths

def manifest_and_story(capture,sources,summary,progress,cases_list,snapshots,quality,stats,roles,api):
    commit=subprocess.check_output(['git','rev-parse','HEAD'],cwd=ROOT,text=True).strip();dirty=bool(subprocess.check_output(['git','status','--porcelain','--','.'],cwd=ROOT,text=True).strip());hashes={}
    for path in ['backend/db.py','backend/collectors.py','backend/service.py','backend/policy_delivery.py','backend/policy_delivery_comparison.py','backend/project_collection.py','backend/stage_evidence.py','backend/evidence_bridge.py','backend/source_operations.py','backend/area_snapshot.py','web/assets/policy-907.js','web/assets/policy-stage-flow.js','web/assets/requirements-manual.js','web/assets/progress-report.js','web/index.html','scripts/extract_presentation_evidence_20261022.py','scripts/package_presentation_evidence_20261022.py']:
        hashes[path]=hashlib.sha256((ROOT/path).read_bytes()).hexdigest()
    dump('code_version.json',dict(commit=commit,working_tree_dirty=dirty,scope='PoC/10-home',files=hashes,note='커밋만으로 실행코드가 같다고 주장하지 않음. 파일해시가 실제 미커밋 작업트리 기준.'))
    copy_text('requirements.txt','12_technical/requirements.txt')
    current=next(r for r in summary if r['stage']=='CONSTRUCTION');first=datetime.fromisoformat(capture['collection_first_at']).astimezone(ZoneInfo('Asia/Seoul')).isoformat();asof=datetime.fromisoformat(capture['extracted_at']).astimezone(ZoneInfo('Asia/Seoul')).isoformat()
    md('00_manifest.md',f'''
# the주택 발표 제작용 근거자료

**서비스 정식 표기: the주택 · 주택공급 디지털 상황판**. 설명상 “더주택”으로 읽는다. PoC10의 공개원천 기반 실증 데이터/화면이며 정부가 운영하는 공식 상황판이 아니다.

- 기준일 **2026-10-22**, 읽기 전용 DB 고정 시각 **{asof}**.
- 대상: **2025-09-07 주택공급 확대방안**, 수도권2026~2030 착공목표135만호 및 공개된 연결사업. 통합 후속대책150만호 전체·전국 모든 사업과 구분한다.
- DB에 남은 최초 수집배치 시작: **{first}**. 팀의 실제 개발 착수일·인력투입시간은 미확인이다.
- 버전: HEAD `{commit[:12]}` + 미커밋 작업트리. 파일별 SHA는 code_version.json. 추출은 REPEATABLE READ READ ONLY로 수행했으며 원천 수집/매핑 변경을 실행하지 않았다.
- 제공 역할: 수집·데이터/백엔드·서비스·검증 담당 자료의 통합 추출. 개인 이름·연락처·인증값·내부 접속주소를 공개자료에서 제거한다.

|구분|수량·단위|
|---|---:|
|등록 정책 / 전체 프로그램|{capture['policy_count']}개 / {capture['program_count']}개|
|9·7 프로그램 / 연결 사업 원천|{capture['policy_program_count']}개 / {capture['linked_source_records']}건|
|전체 사업 원천 /9·7 관리사업|{capture['project_source_records']}건 / **{capture['management_cohorts']}개**|
|9·7 집계에 기여하는 중복 없는 원천ID|{capture['contributing_source_ids']}개|
|등록 수집경로 /성공 이력 경로|{capture['source_count']}개 / {capture['sources_with_any_success']}개|
|현재 공고 원천ID /문서근거ID|{capture['notice_records']:,}건 / {capture['document_records']:,}건|
|목록까지 포함한 문서 inventory|{capture['document_inventory_records']:,}기록|
|현재 저장 레코드(원천·관계·관리 포함)|{capture['current_records']:,}건 — 주택/사업수가 아님|
|고유 보관응답 해시 /manifest 관측|{capture['unique_response_hashes']:,}개 / {capture['manifest_observations']:,}회 — API응답 포함|
|관계 추출 /고정 독립스냅샷|{capture['relations']:,}행 / {capture['metric_snapshot_count']:,}개|

**즉시 쓸 수 있는 핵심:** 공식4분기32행 전부 사업 대응,31개 비교단위 중26개 원수량 동등. 검단 국민991+영구332=1,323, 부평 동일현장700→350, 연남8호 독립사업 확인 등 실제 사례6개를 제공한다.

**발표 전 수정할 문장:** “56개100% 일치”는 입증되지 않았다. 기존53+보충3 목록을 재구성해 각 판정을 제공한다. 공식 원문은99,800호 정확값이 아니라 **9.9만호 이상**으로 표현한다. 현재 **{current['current_units']:,}호는 정책 연결 누적**이며 정부의 올해 수도권 전체 착공과 동일 모집단이 아니다.

**구현:** 수집/파싱·판본·연결검토·지도/시계열·보고서·CSV. **미구현/미확인:** 스캔PDF 통합OCR,실제3종 독립서비스/내외망 분리운영,개인 맞춤 자격판정,전체정부실적 전량대사. 일부확인은사용자위임에따른에이전트담당자대행실DB기록이며실제행정기관담당자의인적결재증명은아니다.

빠른 시작: [제작자용 요약](00_발표제작자_빠른읽기.md) → [56/185/32 대조 엑셀](07_official_validation/07_56_metrics_comparison.xlsx) → [차이 설명](07_official_validation/reconciliation_notes.md) → [제공 현황](00_request_response_matrix.csv). 원본메타·원장/쿼리·스키마는각폴더와12_technical에있다.
''')
    matrix=[]
    matrix_data=[('공급 단계·출처 맵','있음','01_process_sources/01_supply_process_and_sources.xlsx','7단계내부관리정의·76경로·발견/파싱/주기','실제 법정절차 일률화 아님'),('공고·파일 통계','일부','02_collection_starmap/source_stats.csv','전수record/manifest메타·문서목록·스타맵','게시판고유수/스캔분류미측정'),('파이프라인·실행 로그','있음','03_pipeline/pipeline_runs.csv','전체배치 및감사표본/큐','원문별전처리전환율미측정'),('연결·담당자 확인','있음','04_entity_linking/relations.csv','현재관계·검토표본·역할기록','실제공무원인적승인증명미확인'),('정제·중복/변경 사례','있음','05_case_studies/case_index.csv','실제사례6개','OCR성공/법정분할·합병사례는없음'),('여러 시점 데이터','일부','06_data_quality/dataset_snapshots.csv','10/12·10/21·10/22 보관자료','초기전과정동일형식스냅샷미보유'),('공식56개 비교','일부','07_official_validation/07_56_metrics_comparison.xlsx','53+보충3재구성·185원표·32사업·8사례·10통계','56전부일치근거미보유'),('99,800 vs82,968','일부','07_official_validation/reconciliation_notes.md','정책누적재현 및공식9.9만이상 확인','99,800정확값원장/동일범위분해미확보'),('3액터 화면·권한','일부','08_service/08_roles_and_permissions.xlsx','현재PoC관점별화면·관리자저장경계','실제3서비스·내외망분리미확인'),('지도·사업·단계원장','있음','09_map_dashboard/project_map.csv','177관리사업·단계근거/계획·상태이력','추정/계획/실제분리'),('팀 캐릭터·동영상','없음','10_story_assets/character_assets/README.md','보유/사용허가미확인','새가상캐릭터/인력실적은생성하지않음'),('총리 반복발언 연혁','일부','11_official_mission/mission_sources.csv','10.16공식상황판계획과총리점검체계자료','반복동일요구발언별원문전수미확보')]
    for name,held,path,scope,limit in matrix_data:matrix.append(dict(항목=name,보유여부=held,실제파일=path,수집기준일='2026-10-22 / 원작성시점별도',제공가능범위=scope,구현여부='파일별구현/미구현구분',비고=limit))
    csvout('00_request_response_matrix.csv',matrix)
    md('00_request_response_matrix.md','# 요청 항목 회신\n\n'+table(['항목','보유','파일','제공 범위','한계'],[[r['항목'],r['보유여부'],r['실제파일'],r['제공가능범위'],r['비고']] for r in matrix]))
    md('00_발표제작자_빠른읽기.md',f'''
# 발표 제작자는 여기부터 읽으세요

## 사용할 수 있는 서사와 숫자

1. **미션:** 10.16공식발표에서디지털주택공급상황판계획과상세정보·계획·실적공개방향을확인했다. 총리의반복발언별전수원문은미확보이므로반복된동일요구장면을사실처럼새로만들지않는다.
2. **현실의분산:** 등록수집경로{capture['source_count']}개,기관표기31개,대표도메인22개. 기관표기와조직코드상기관수는다르다. API응답파일까지포함한고유해시{capture['unique_response_hashes']:,}개를‘공고문16만개’라고표현하지않는다.
3. **돌파:** 기관별필드·원천ID·블록·유형을정규화하고원문·판본·확인상태를연결했다. 스캔PDF전체자동OCR이아니라텍스트PDF/HTML/HWPX등자료형식별파서와수동보완의조합이다.
4. **실증사례:** 연남8/연희96구분,부평700→350중복정리,검단국민991+영구332구성,대장A9 568→574판본,고덕3일정/범위보정,왕숙2 1063/1082미해결충돌보존. 각사례README의6단계구성을그대로슬라이드근거로쓸수있다.
5. **데이터모델:** 전체사업원천452→9·7연결232→집계기여원천221→관리사업177. 차이는자료를잃어버린것이아니라정책범위·집계채택·동일현장중복을구분한분모다.
6. **검증:** 4분기32행사업명은전부대응,중복정리31비교단위중26원수량동등. 나머지5건은병점범위차이1건확인·나머지원인미확정. 총계+22의상쇄를100%정확도로바꾸지않는다.
7. **기능:** 지도중심정책/프로그램/사업조회·현황보고서·근거추적·지역여건은실제화면으로보여준다. 세사용자관점과실제3개독립서비스/내외망분리운영은구분한다.

## 숫자 간 섞으면 안 되는 것

- 누적착공 **{current['current_units']:,}호**,연말누적예상 **{current['year_end_units']:,}호**는177관리사업의정책누적이다. 이전연도착공과채택추정이포함된다.
- 정부당해연도수도권전체 **9.9만호이상**과같은모집단이아니다.99,800이라는정확값은근거미확보다.
- 기존53+보충3=56개명단을제공하지만 **56개100%일치**는사용할수없다.185항목원표도‘전체정확도’가아니다.
- 공식통계수집값39,940등을개별사업합산으로독립검증한숫자처럼설명하지않는다.
- 현재 저장 레코드 469만 건에는 원천·관계·관리 기록이 포함된다. 이를 469만 개 공고·사업·가구라고 말하지 않는다.

## 파일 읽는 순서

`00_manifest.md` → `07_official_validation` 엑셀·차이설명 → `05_case_studies` → `02_collection_starmap` → `08_service` 화면·구성도 → `09_map_dashboard` 집계표. 큰00_data·전수파일inventory·상태이력은검산할때만열면된다.

이패키지는추출자료이며슬라이드완성본이아니다. 수치·확인상태·미해결범위를그대로사용하면외부질문에도근거파일로답할수있다.
''')
    mission=[dict(id='MISSION-01',official_date='2026-10-16',source='https://www.korea.kr/briefing/policyBriefingView.do?newsId=156785506',speaker_role='국무조정실·국토교통부 정례브리핑',confirmed='디지털주택공급상황판 구축·운영 계획,상세단계정보/향후계획/실적 공개 방향',implemented_mapping='정책/프로그램/사업 지도·단계·시계열·원문·보고서',limit='정부의공식운영시스템납품또는실제사용검증이아님'),dict(id='MISSION-02',official_date='2026-10-16',source='https://www.fsc.go.kr/po010101/87948',speaker_role='관계부처 공동 보도자료',confirmed='8.14총리주재점검체계출범과누적3회회의·6회현장방문설명',implemented_mapping='현황점검·근거사업드릴다운·병목/조치관리 기능',limit='회의별상황판요구발언원문전수는미확보')]
    csvout('11_official_mission/mission_sources.csv',mission)
    md('11_official_mission/README.md','''# 공식 미션 근거

mission_sources.csv의두공식URL을2026-10-22확인했다.디지털상황판계획및총리점검체계는확인할수있다.다만여러회의에서동일한상황판요구가반복됐다는날짜별발언목록은이번제출에서확보하지못했다.프로젝트의사용자제공요구사항61개는그자체로공식총리발언원문이아니다.
''')
    work=[dict(event_date='2026-10-12',event='요구사항·초기데이터범위점검',evidence='references/requirements_20261012_snapshot.json',measured='원천사업292건',note='최초크롤링완료시점이라고단정하지않음'),dict(event_date='2026-10-20',event='동일현장중복·연남독립사업확인',evidence='05_case_studies',measured='부평700→350;연남8호별도',note='인력투입시간미측정'),dict(event_date='2026-10-21',event='읽기전용풀데이터·185공식항목대조',evidence='references/snapshot_20261021.json',measured='원천452·연결231·관리176',note='모든185개일치아님'),dict(event_date='2026-10-22',event='누적·연말산식·공식비교성적표보완',evidence='06_data_quality/scorecard_latest.md',measured='관리177·누적착공82968·연말99885',note='연간정부전체착공과다름'),dict(event_date='2026-10-22',event='보고서·수집항목·지역표시자료사전계산',evidence='08_service/screenshots',measured='실제UI 캡처',note='외부망운영성과로확대하지않음')]
    csvout('10_story_assets/work_log.csv',work);md('10_story_assets/character_assets/README.md','# 팀 캐릭터·동영상\n\n사용허가가확보된팀캐릭터원본/포즈세트와팀원시연동영상은이번추출범위에서보유여부를확인하지못했다.임의캐릭터나가상작업성과는생성하지않았다.실제서비스화면캡처는08_service/screenshots에제공한다.')
    return matrix

def main():
    capture=read('00_data/capture_summary.json');projects=read('00_data/projects.json');programs=read('00_data/programs.json');sources=read('00_data/current_sources.json');model=read('00_data/policy_model.json')
    summary,maps,events,targets,ledger=identity_and_map(model,projects,programs)
    assert next(r['current_units'] for r in summary if r['stage']=='CONSTRUCTION')==82968
    roster,detailed,quantity,starts,statistics=official_validation(summary)
    definitions,registry,progress=process_and_pipeline(capture,sources,programs,projects,summary)
    case_rows=cases(projects,quantity,starts)
    snapshots,q,issues=quality_and_snapshots(capture,projects,model,summary,ledger,detailed,quantity)
    stats=starmap(capture,sources,projects,programs)
    roles,api=system_and_dictionary()
    matrix=manifest_and_story(capture,sources,summary,progress,case_rows,snapshots,q,stats,roles,api)
    # A complete scoped stage export can be read independently of the 452-source inventory.
    from backend.collection_exports import export
    from types import SimpleNamespace
    (OUT/'09_map_dashboard/full_source_stage_dataset.csv').write_text(export(SimpleNamespace(projects=lambda **_:projects),read('00_data/collection_matrix.json')))
    workbook('00_발표핵심자료.xlsx',[('주요수치',[dict(item=k,value=v) for k,v in capture.items() if not isinstance(v,(dict,list))]),('현재단계집계',summary),('요청회신',matrix),('진행률정의',progress),('사례목록',case_rows),('과거현재스냅샷',snapshots),('이번정합성검사',q),('원천분산도',stats),('공식32사업',quantity),('공식착공8사례',starts),('3액터구분',roles)])
    dump('package_build_summary.json',dict(built_at=datetime.now(ZoneInfo('Asia/Seoul')).isoformat(),capture=capture,stages=summary,case_count=len(case_rows),official_roster_count=len(roster),official_comparison_rows=len(detailed),current_checks=q,notes='이자료는실측추출·보관검증을구분한다.공식56전부일치로표현하지않음.'))
    print(json.dumps(dict(manifest=str(OUT/'00_manifest.md'),core_xlsx=str(OUT/'00_발표핵심자료.xlsx'),stage_summary=summary,checks=q),ensure_ascii=False),flush=True)

if __name__=='__main__':main()
