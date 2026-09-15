from __future__ import annotations

import io
import os
import re
import subprocess
import tempfile
import zipfile
from dataclasses import dataclass
from pathlib import Path, PurePosixPath
from xml.etree import ElementTree

from docx import Document
from pypdf import PdfReader

from ..config import get_settings
from .filename import SAFE_EXTENSIONS


OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
ZIP_MAGIC = b"PK\x03\x04"


@dataclass(frozen=True)
class ParseResult:
    status: str
    text: str = ""
    error: str | None = None


def is_safe_zip_member(name: str) -> bool:
    normalized = name.replace("\\", "/")
    path = PurePosixPath(normalized)
    return bool(normalized and not path.is_absolute() and ".." not in path.parts and not re.match(r"^[A-Za-z]:", normalized))


def _decode_text(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp949", "euc-kr"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _parse_pdf(content: bytes) -> str:
    return "\n\n".join((page.extract_text() or "") for page in PdfReader(io.BytesIO(content)).pages)


def _parse_docx(content: bytes) -> str:
    document = Document(io.BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(paragraphs)


def _parse_hwpx(content: bytes) -> str:
    chunks: list[str] = []
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        for info in archive.infolist():
            if not is_safe_zip_member(info.filename):
                raise ValueError(f"안전하지 않은 HWPX 경로: {info.filename}")
            if not info.filename.startswith("Contents/") or not info.filename.endswith(".xml"):
                continue
            root = ElementTree.fromstring(archive.read(info))
            chunks.extend(element.text or "" for element in root.iter() if element.tag.endswith("}t"))
    return "\n".join(chunk for chunk in chunks if chunk.strip())


def _parse_xlsx(content: bytes, max_uncompressed_bytes: int = 100 * 1024 * 1024) -> str:
    """Extract visible cell values from OOXML without executing formulas/macros."""
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        infos = archive.infolist()
        if any(not is_safe_zip_member(info.filename) for info in infos):
            raise ValueError("안전하지 않은 XLSX 내부 경로")
        if sum(info.file_size for info in infos) > max_uncompressed_bytes:
            raise ValueError("XLSX 압축 해제 크기 제한 초과")
        names = {info.filename for info in infos}
        shared: list[str] = []
        if "xl/sharedStrings.xml" in names:
            root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
            for item in root.iter():
                if item.tag.endswith("}si"):
                    shared.append("".join(
                        child.text or "" for child in item.iter() if child.tag.endswith("}t")
                    ))
        lines: list[str] = []
        sheets = sorted(name for name in names if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name))
        for sheet_name in sheets:
            root = ElementTree.fromstring(archive.read(sheet_name))
            for row in (element for element in root.iter() if element.tag.endswith("}row")):
                values: list[str] = []
                for cell in (element for element in row if element.tag.endswith("}c")):
                    cell_type = cell.attrib.get("t", "")
                    value_node = next((node for node in cell.iter() if node.tag.endswith("}v")), None)
                    if cell_type == "inlineStr":
                        value = "".join(
                            node.text or "" for node in cell.iter() if node.tag.endswith("}t")
                        )
                    elif value_node is None or value_node.text is None:
                        value = ""
                    elif cell_type == "s":
                        index = int(value_node.text)
                        value = shared[index] if 0 <= index < len(shared) else ""
                    else:
                        value = value_node.text
                    if value.strip():
                        values.append(value.strip())
                if values:
                    lines.append(" | ".join(values))
        return "\n".join(lines)


def _parse_hwp_cli(content: bytes) -> str:
    settings = get_settings()
    max_output = settings.hwp_parse_max_output_mb * 1024 * 1024
    with tempfile.NamedTemporaryFile(suffix=".hwp") as source:
        source.write(content)
        source.flush()
        completed = subprocess.run(
            [settings.hwp_cli_path, "cat", "--format", "markdown", source.name],
            capture_output=True,
            check=False,
            timeout=settings.hwp_parse_timeout_seconds,
            env={**os.environ, "RUST_BACKTRACE": "0"},
            start_new_session=True,
        )
    if completed.returncode:
        detail = completed.stderr.decode("utf-8", errors="replace").strip()
        raise RuntimeError(f"hwp-cli 종료 코드 {completed.returncode}: {detail[:400]}")
    if len(completed.stdout) > max_output:
        raise ValueError(f"HWP 추출 본문이 {settings.hwp_parse_max_output_mb}MB 제한을 초과함")
    return completed.stdout.decode("utf-8", errors="replace")


def _parse_zip(content: bytes, max_uncompressed_bytes: int = 100 * 1024 * 1024) -> str:
    chunks: list[str] = []
    total = 0
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        for info in archive.infolist():
            if not is_safe_zip_member(info.filename):
                raise ValueError(f"ZIP path traversal 차단: {info.filename}")
            total += info.file_size
            if total > max_uncompressed_bytes:
                raise ValueError("ZIP 압축 해제 크기 제한 초과")
            ext = Path(info.filename).suffix.lower().lstrip(".")
            if ext not in SAFE_EXTENSIONS or ext == "zip" or info.is_dir():
                continue
            child = parse_bytes(archive.read(info), ext)
            if child.status == "parsed" and child.text.strip():
                chunks.append(f"\n--- {Path(info.filename).name} ---\n{child.text}")
    return "\n".join(chunks)


def parse_bytes(content: bytes, extension: str) -> ParseResult:
    ext = extension.lower().lstrip(".")
    if not content:
        return ParseResult("failed", error="0바이트 파일")
    try:
        if ext == "txt":
            text = _decode_text(content)
        elif ext == "pdf":
            text = _parse_pdf(content)
        elif ext == "docx":
            text = _parse_docx(content)
        elif ext == "xlsx":
            text = _parse_xlsx(content)
        elif ext in {"hwp", "hwpx"}:
            # 나라장터에는 확장자만 HWPX이고 실제 본문은 OLE HWP인 파일도 있다.
            if content.startswith(OLE_MAGIC):
                text = _parse_hwp_cli(content)
            elif content.startswith(ZIP_MAGIC) or zipfile.is_zipfile(io.BytesIO(content)):
                text = _parse_hwpx(content)
            else:
                raise ValueError(f"{ext.upper()} 파일 시그니처가 올바르지 않음")
        elif ext == "zip":
            text = _parse_zip(content)
        else:
            return ParseResult("unsupported", error=f"지원하지 않는 형식: {ext or 'unknown'}")
        cleaned = text.replace("\x00", "").strip()
        return ParseResult("parsed", cleaned) if cleaned else ParseResult("failed", error="추출된 본문이 없음")
    except Exception as exc:
        return ParseResult("failed", error=f"{type(exc).__name__}: {str(exc)[:500]}")


def parse_file(path: Path, extension: str | None = None) -> ParseResult:
    return parse_bytes(path.read_bytes(), extension or path.suffix)
